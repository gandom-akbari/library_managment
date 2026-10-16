from openpyxl import Workbook, load_workbook
from datetime import date, timedelta
import os

FILE = "membership.xlsx"
LOAN_DAYS = 14


def open_workbook():
    if os.path.exists(FILE):
        wb = load_workbook(FILE)
    else:
        wb = Workbook()
        wb.active.title = "membership"
        wb["membership"].append(["id", "name", "age"])
    if "loans" not in wb.sheetnames:
        ws = wb.create_sheet("loans")
        ws.append(["member_id", "book", "loan_date", "return_date", "returned", "returned_date"])
    return wb


def member_ids(wb):
    ids = []
    for row in wb["membership"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            ids.append(row[0])
    return ids


def lend_book(wb):
    try:
        member_id = int(input("member id: "))
    except ValueError:
        print("id must be a number")
        return

    if member_id not in member_ids(wb):
        print("no member with this id")
        return

    book = input("book name: ").strip()
    if not book:
        print("book name can't be empty")
        return

    loan_date = date.today()
    return_date = loan_date + timedelta(days=LOAN_DAYS)

    wb["loans"].append([member_id, book, str(loan_date), str(return_date), "no", ""])
    print(f"lent '{book}' to member {member_id}. return before {return_date}")


if __name__ == "__main__":
    wb = open_workbook()

    lend_book(wb)

    wb.save(FILE)
    print("saved at:", os.path.abspath(FILE))