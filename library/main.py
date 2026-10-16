from openpyxl import Workbook, load_workbook
from datetime import date, timedelta
import os

FILE = "membership.xlsx"

wb = load_workbook(FILE) if os.path.exists(FILE) else Workbook()
if "membership" not in wb.sheetnames:
    if wb.sheetnames == ["Sheet"]:
        wb.active.title = "membership"
    else:
        wb.create_sheet("membership")
    wb["membership"].append(["id", "name", "age"])
if "books" not in wb.sheetnames:
    wb.create_sheet("books").append(["id", "title"])
if "loans" not in wb.sheetnames:
    wb.create_sheet("loans").append(
        ["member_id", "book", "loan_date", "return_date", "returned", "returned_date"]
    )


def rows(sheet):
    return list(wb[sheet].iter_rows(min_row=2))


def ask_int(text):
    try:
        return int(input(text))
    except ValueError:
        print("must be a number")
        return None


def add_member():
    name = input("name: ")
    age = ask_int("age: ")
    if age is None:
        return
    new_id = len(rows("membership")) + 1
    wb["membership"].append([new_id, name, age])
    print("member added with id", new_id)


def add_book():
    title = input("book name: ")
    wb["books"].append([len(rows("books")) + 1, title])
    print("book added")


def lend_book():
    member_id = ask_int("member id: ")
    if member_id not in [r[0].value for r in rows("membership")]:
        print("no member with this id")
        return
    book = input("book name: ")
    if book not in [r[1].value for r in rows("books")]:
        print("add the book first")
        return
    today = date.today()
    due = today + timedelta(days=14)
    wb["loans"].append([member_id, book, str(today), str(due), "no", ""])
    print("lent. return before", due)


def return_book():
    member_id = ask_int("member id: ")
    active = [r for r in rows("loans") if r[0].value == member_id and r[4].value == "no"]
    if not active:
        print("this member has no borrowed books")
        return
    print("borrowed books:")
    for r in active:
        print("-", r[1].value, "(due", r[3].value + ")")
    book = input("which book is returned? ")
    for r in active:
        if r[1].value == book:
            r[4].value = "yes"
            r[5].value = str(date.today())
            print("OK, returned")
            return
    print("not found")


while True:
    print()
    print("1) add member")
    print("2) add book")
    print("3) lend book")
    print("4) return book")
    print("5) save and exit")
    choice = input("choose: ")

    if choice == "1":
        add_member()
    elif choice == "2":
        add_book()
    elif choice == "3":
        lend_book()
    elif choice == "4":
        return_book()
    elif choice == "5":
        break
    else:
        print("invalid choice")

wb.save(FILE)
print("saved at:", os.path.abspath(FILE))