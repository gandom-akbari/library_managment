from library.member_class import Member
from openpyxl import Workbook, load_workbook
import os

FILE = "membership.xlsx"


def open_workbook():
    if os.path.exists(FILE):
        wb = load_workbook(FILE)
    else:
        wb = Workbook()
        wb.active.title = "membership"
        wb["membership"].append(["id", "name", "age"])
    if "loans" not in wb.sheetnames:
        ws = wb.create_sheet("loans")
        ws.append(["member_id", "book", "loan_date", "return_date", "returned", "returned_date"])
    return wb


def load_members(wb):
    ws = wb["membership"]
    members = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        member_id, name, age = row
        members.append(Member(name, age, member_id))
    return members


def save(wb):
    wb.save(FILE)
    print("saved at:", os.path.abspath(FILE))