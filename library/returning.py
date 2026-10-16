from openpyxl import Workbook, load_workbook
from datetime import date
import os

FILE = "membership.xlsx"


def open_workbook():
    if os.path.exists(FILE):
        wb = load_workbook(FILE)
    else:
        wb = Workbook()
        wb.active.title = "membership"
        wb["membership"].append(["id", "name", "age"])
    if "loans" not in wb.sheetnames:
        ws = wb.create_sheet("loans")
        ws.append(["member_id", "book", "loan_date", "return_date", "returned", "returned_date"])
    return wb


def return_book(wb):
    ws = wb["loans"]

    try:
        member_id = int(input("member id: "))
    except ValueError:
        print("id must be a number")
        return

    # پیدا کردن کتاب‌هایی که این عضو قرض گرفته و هنوز برنگردانده
    active = []
    for row_number in range(2, ws.max_row + 1):
        if (
            ws.cell(row=row_number, column=1).value == member_id
            and ws.cell(row=row_number, column=5).value == "no"
        ):
            active.append(row_number)

    if not active:
        print("this member has no borrowed books")
        return

    print("borrowed books:")
    for number, row_number in enumerate(active, start=1):
        book = ws.cell(row=row_number, column=2).value
        due = ws.cell(row=row_number, column=4).value
        print(f"{number}) {book} (due {due})")

    # اگر فقط یک کتاب باشد، خودش همان را برمی‌دارد
    if len(active) == 1:
        chosen = 1
    else:
        try:
            chosen = int(input("which one is returned? (number): "))
        except ValueError:
            print("enter a number")
            return
        if chosen < 1 or chosen > len(active):
            print("invalid number")
            return

    row_number = active[chosen - 1]
    book = ws.cell(row=row_number, column=2).value

    ws.cell(row=row_number, column=5).value = "yes"
    ws.cell(row=row_number, column=6).value = str(date.today())

    due = date.fromisoformat(ws.cell(row=row_number, column=4).value)
    late_days = (date.today() - due).days

    print(f"OK, '{book}' returned")
    if late_days > 0:
        print(f"it was {late_days} days late")


